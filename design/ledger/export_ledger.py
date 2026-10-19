"""Read-only CSV export of PROJECT_A_OPEN_ITEMS_LEDGER.xlsx (derived; the workbook is canonical).

Run from the repository root:  python design/ledger/export_ledger.py
Writes design/ledger/export/{ledger,archive,questions,links,aliases,change-log}.csv (UTF-8, LF).
Formula columns (Active?, Review flag, Unanswered linked questions, Answered?, Needs review?, counts, alias keys)
are computed here with the same logic as the workbook formulas, so the export holds values even when the
workbook has not been recalculated. Nothing is written back to the workbook.
"""
import csv, os, sys
from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, "PROJECT_A_OPEN_ITEMS_LEDGER.xlsx")
OUT = os.path.join(HERE, "export")

def table(ws, key_col=1, stop_blank=True):
    rows = [[c.value for c in r] for r in ws.iter_rows()]
    hdr = rows[0]
    data = []
    for r in rows[1:]:
        if r[key_col - 1] in (None, ""):
            continue
        data.append(r)
    return hdr, data

def clean(v):
    if v is None:
        return ""
    if isinstance(v, str) and v.startswith("="):
        return ""
    return v

def write(name, hdr, rows):
    os.makedirs(OUT, exist_ok=True)
    with open(os.path.join(OUT, name), "w", encoding="utf-8", newline="") as f:
        w = csv.writer(f, lineterminator="\n")
        w.writerow(hdr)
        for r in rows:
            w.writerow(r)

def main():
    wb = load_workbook(SRC)  # formulas as text; values for formula cells are computed below
    # vocabularies
    lists = wb["Lists"]
    active_map = {}
    for r in range(2, lists.max_row + 1):
        s, a = lists.cell(row=r, column=1).value, lists.cell(row=r, column=2).value
        if s: active_map[s] = a
    qh, qrows = table(wb["Questions"])
    qi = {h: i for i, h in enumerate(qh)}
    answered = {r[0]: bool(r[qi["YOUR ANSWER (type here)"]]) for r in qrows}
    reviewed = {r[0]: r[qi["Review status"]] == "REVIEWED" for r in qrows}
    lh, lrows = table(wb["Links"])
    links = []
    for r in lrows:
        qid, lid, eff = r[0], r[1], r[2]
        ans = answered.get(qid)
        links.append((qid, lid, eff, "?" if ans is None else ("Yes" if ans else "No"),
                      "Yes" if (ans and not reviewed.get(qid, False)) else "No"))
    # ledger / archive
    def ledger_like(ws, archive):
        h, rows = table(ws)
        ix = {k: i for i, k in enumerate(h)}
        out = []
        for r in rows:
            r = list(r)
            lid = r[0]
            if archive:
                r[ix["Active?"]], r[ix["Review flag"]], r[ix["Unanswered linked questions"]] = "No", "", ""
            else:
                r[ix["Active?"]] = active_map.get(r[ix["Current status"]], "?")
                r[ix["Review flag"]] = "REVISE" if any(l[1] == lid and l[4] == "Yes" for l in links) else ""
                r[ix["Unanswered linked questions"]] = sum(1 for l in links if l[1] == lid and l[3] == "No" and l[2] != "INFORMS")
            out.append([clean(v) for v in r])
        return h, out
    h, rows = ledger_like(wb["Ledger"], False); write("ledger.csv", h, rows); nled = len(rows)
    h, rows = ledger_like(wb["Archive"], True); rows = [r for r in rows if " " not in str(r[0])]; write("archive.csv", h, rows); narc = len(rows)
    # questions
    out = []
    for r in qrows:
        r = list(r)
        r[qi["Ledger items open because of it (formula)"]] = sum(1 for l in links if l[0] == r[0] and l[2] != "INFORMS")
        out.append([clean(v) for v in r])
    write("questions.csv", qh, out)
    write("links.csv", ["QID", "LedgerID", "Effect", "Answered?", "Needs review?"], links)
    ah, arows = table(wb["Aliases"])
    write("aliases.csv", ah, [[clean(r[0]), clean(r[1]), clean(r[2]), clean(r[3]), f"{r[0]}|{r[2]}"] for r in arows])
    ch, crows = table(wb["Change log"])
    crows = [r for r in crows if not (isinstance(r[0], str) and r[0].startswith("Append-only"))]
    write("change-log.csv", ch, [[clean(v) for v in r] for r in crows])
    print(f"exported: ledger {nled} rows, archive {narc}, questions {len(qrows)}, links {len(links)}, aliases {len(arows)}, change log {len(crows)}")

if __name__ == "__main__":
    sys.exit(main())
