"""Record the human's API probe (INT-T-009, REQ-U22-11) in the ledger, in place (never rebuilt): Q-U33 answered, N-31, C-X07 and
C-S03/DEFUSAGE updated, new row N-37 for what the probe left open. Appends Change log rows above the append-only note. Refuses to
run twice. Run from the repository root, then run design/ledger/export_ledger.py.
"""
import re, sys
from openpyxl import load_workbook

SRC = "design/ledger/PROJECT_A_OPEN_ITEMS_LEDGER.xlsx"
D, BY = "2026-10-27", "Claude"
SRCN = "INT-T-009 (requirements/intake/2026-10-27-user-report-api-probe-run.md); design/package-v5/10-p0-university-api-reading.md"
wb = load_workbook(SRC)
L, Q, C = wb["Ledger"], wb["Questions"], wb["Change log"]
lh = [c.value for c in L[1]]
qh = [c.value for c in Q[1]]
rows = {L.cell(r, 1).value: r for r in range(2, L.max_row + 1) if L.cell(r, 1).value}
qrows = {Q.cell(r, 1).value: r for r in range(2, Q.max_row + 1) if Q.cell(r, 1).value}
if "N-37" in rows:
    sys.exit("already applied; nothing written")
col = lambda n: lh.index(n) + 1
qc = lambda n: qh.index(n) + 1
log = []


def note(lid, field, old, new, reason):
    log.append([D, lid, field, old, new, reason, "INT-T-009 / V5", BY])


def setf(lid, field, new, reason):
    r = rows[lid]; c = col(field); old = L.cell(r, c).value; L.cell(r, c).value = new; note(lid, field, old, new, reason)


def known(lid, text):
    r = rows[lid]; c = col("What is known now"); old = L.cell(r, c).value or ""
    L.cell(r, c).value = (old + " | " if old else "") + f"V5 {D}: " + text
    note(lid, "What is known now", "(appended)", text[:180], "design/package-v5/10")


# Q-U33 answered by action
r = qrows["Q-U33"]
for k, v in {"YOUR ANSWER (type here)": "[Done] INT-T-009 line 1: the human ran the probe on gpt-oss-120b and five further models; outputs in design/package-v5/inputs/ (first round); a second probe (p0_api_probe_v2.py) is still to be run",
             "Answer type": "Document provided", "Answer date": D, "Review status": "NEEDS FOLLOW-UP", "Provenance (verbatim? where from)": SRCN,
             "Revised on / by": f"{D} / {BY}", "Resulting ledger changes": "see the Change log of 2026-10-27"}.items():
    Q.cell(r, qc(k)).value = v
note("Q-U33", "Answer", "", "answered by action (REQ-U22-11)", SRCN)

SUM = "six models: cache counters (cached_tokens, created_cache_tokens) and reasoning_tokens reported; prefix cache partial and block-wise (59-99.7% on a repeat), kept on an appended turn; allowed_tools rejected (HTTP 400, LiteLLM proxy); GLM-5.3 limited to 2 requests; tool calling not shown; no prices"
known("N-31", SUM)
setf("N-31", "Next action", "Second probe: the human runs design/package-v5/tools/p0_api_probe_v2.py (changed tools array, tool calling at max_tokens 400, an absent tool)", "design/package-v5/10 s6")
setf("N-31", "Status date", D, "INT-T-009")
known("C-X07", "cache fields exist on the university API (first round of N-31); the billed price is not reported, so the bill is computed from counts and a named price table (DPR-0006 s1)")
setf("C-X07", "Status date", D, "INT-T-009")
known("C-S03/DEFUSAGE", "on the university route vendor deferral does not exist; a changed tools array and appended definitions are tested by the second probe (C3, C4)")
setf("C-S03/DEFUSAGE", "Status date", D, "INT-T-009")

# new row N-37
tmpl = rows["N-30"]; last = max(rows.values()); r = last + 1
for cname in ("Active?", "Review flag", "Unanswered linked questions"):
    f = L.cell(tmpl, col(cname)).value
    L.cell(r, col(cname)).value = re.sub(rf"(?<![0-9$]){tmpl}(?![0-9])", str(r), f)
vals = {"LedgerID": "N-37", "Title": "University API: tool calling, tool_choice and a changed tools array (second probe)", "Current status": "OPEN",
        "Closes by (primary)": "TEST-PILOT", "Cluster": "K9 Cache & billing", "Must precede (gate)": "G2 Before the affected pilot/arm",
        "Priority": "Normal", "Owner": "Lead; you run the probe", "Next action": "The human runs p0_api_probe_v2.py (GLM-5.3 with --pause 35)",
        "Kind": "Test", "Level": "Item",
        "What is open (description)": "Whether the models call tools at all through the university API, whether tool_choice required/named is honoured (MiniMax-M3.1-Flash ignored required), what a changed tools array does to the cache (V5-R02 on this route), and how models answer a request for a tool that does not exist",
        "What is known now": "First round (2026-10-27) showed no tool call on any model, all inside a 16-token cap (design/package-v5/10)",
        "Closed when": "The second probe's output is read and its limits stated", "Change vs baseline": "NEW (V5, 2026-10-27)", "Status date": D,
        "Status basis (evidence pointer)": "design/package-v5/10-p0-university-api-reading.md", "Needs new AUTH?": "No (AUTH-020)",
        "Origin (first source)": "design/package-v5/10 s3", "Linked questions": ""}
for k, v in vals.items():
    L.cell(r, col(k)).value = v
note("N-37", "(new row)", "", "OPEN", "design/package-v5/10")

note_row = next(i for i in range(C.max_row, 1, -1) if str(C.cell(i, 1).value or "").startswith("Append-only"))
C.insert_rows(note_row - 1, amount=len(log))
for i, row in enumerate(log):
    for j, v in enumerate(row):
        C.cell(note_row - 1 + i, j + 1).value = v
bad = [i for i in range(2, L.max_row + 1) if L.cell(i, 1).value and any(not (isinstance(L.cell(i, c).value, str) and re.search(rf"{ref}{i}\b", L.cell(i, c).value)) for c, ref in ((4, "C"), (5, "A"), (6, "A")))]
assert not bad, f"formula check failed at rows {bad[:5]}"
wb.save(SRC)
print("change-log rows", len(log))
