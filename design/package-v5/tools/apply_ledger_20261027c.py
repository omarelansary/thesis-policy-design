"""Record the human's instruction of 2026-10-27 (INT-T-011) in the ledger, in place (never rebuilt): N-38 (model choice) DEFERRED with a
reopen trigger; N-39 (price table) PARTLY ANSWERED from design/package-v5/11-price-ranges.md; N-37 note. Appends Change log rows above the
append-only note. Refuses to run twice. Run from the repository root, then run design/ledger/export_ledger.py.
"""
import re, sys
from openpyxl import load_workbook

SRC = "design/ledger/PROJECT_A_OPEN_ITEMS_LEDGER.xlsx"
D, BY = "2026-10-27", "Claude"
wb = load_workbook(SRC)
L, C = wb["Ledger"], wb["Change log"]
lh = [c.value for c in L[1]]
rows = {L.cell(r, 1).value: r for r in range(2, L.max_row + 1) if L.cell(r, 1).value}
col = lambda n: lh.index(n) + 1
if L.cell(rows["N-38"], col("Current status")).value == "DEFERRED":
    sys.exit("already applied; nothing written")
log = []


def note(lid, field, old, new, reason):
    log.append([D, lid, field, old, new, reason, "INT-T-011 / V5", BY])


def setf(lid, field, new, reason):
    r = rows[lid]; c = col(field); old = L.cell(r, c).value; L.cell(r, c).value = new; note(lid, field, old, new, reason)


def known(lid, text):
    r = rows[lid]; c = col("What is known now"); old = L.cell(r, c).value or ""
    L.cell(r, c).value = (old + " | " if old else "") + f"V5 {D}: " + text
    note(lid, "What is known now", "(appended)", text[:180], "design/package-v5/11")


setf("N-38", "Current status", "DEFERRED", "INT-T-011 (REQ-U22-14): model choice is not our case for now; the work must not depend on models")
setf("N-38", "Reopen trigger", "The first experiment is authorized and a run needs a model; or a result shows model dependence", "REQ-U22-14")
setf("N-38", "Next action", "none until the trigger; harnesses take the model as a run-time argument and label every record with it", "REQ-U22-14")
setf("N-38", "Status date", D, "INT-T-011")
setf("N-38", "Status basis (evidence pointer)", "requirements/intake/2026-10-27-user-model-choice-deferred-prices-experiments.md; REQ-U22-14", "INT-T-011")

setf("N-39", "Current status", "PARTLY ANSWERED", "design/package-v5/11-price-ranges.md (REQ-U22-15: the lead searches price sources)")
known("N-39", "price ranges from benchmark and aggregator pages (Aug-Oct 2026, stored): cache-read/input r 0.14 to 1.0 (six of 18 gpt-oss-120b hosts list no cache discount), write fee none listed to 1.25x, output/input 3.1 to 7.5; the break-evens of V5 (D-04) barely depend on these ratios, so results are reported as token classes plus (r, w, o), with dollars under reference cases only when wanted")
setf("N-39", "Next action", "The human accepts or changes the handling in design/package-v5/11 section 7 (token classes plus ratios; dollars under one or two reference cases)", "design/package-v5/11 s7")
setf("N-39", "Status date", D, "INT-T-011")
setf("N-39", "Status basis (evidence pointer)", "design/package-v5/11-price-ranges.md; design/package-v5/calc/price_sensitivity.stdout.txt", "INT-T-011")

known("N-37", "readiness for experiments assessed in design/package-v5/12 (P5 ready once authorized; third probe advised first)")
setf("N-37", "Status date", D, "INT-T-011")

note_row = next(i for i in range(C.max_row, 1, -1) if str(C.cell(i, 1).value or "").startswith("Append-only"))
C.insert_rows(note_row - 1, amount=len(log))
for i, row in enumerate(log):
    for j, v in enumerate(row):
        C.cell(note_row - 1 + i, j + 1).value = v
bad = [i for i in range(2, L.max_row + 1) if L.cell(i, 1).value and any(not (isinstance(L.cell(i, c).value, str) and re.search(rf"{ref}{i}\b", L.cell(i, c).value)) for c, ref in ((4, "C"), (5, "A"), (6, "A")))]
assert not bad, f"formula check failed at rows {bad[:5]}"
wb.save(SRC)
print("change-log rows", len(log))
