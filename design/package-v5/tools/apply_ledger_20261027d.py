"""Record the third probe (INT-T-012, REQ-U22-17) and the literature-verification instruction (REQ-U22-18) in the ledger, in place (never rebuilt):
N-37 and N-31 updated, N-39 note, new row N-40 for the verification review. Appends Change log rows above the append-only note. Refuses to run
twice. Run from the repository root, then run design/ledger/export_ledger.py.
"""
import re, sys
from openpyxl import load_workbook

SRC = "design/ledger/PROJECT_A_OPEN_ITEMS_LEDGER.xlsx"
D, BY = "2026-10-27", "Claude"
wb = load_workbook(SRC)
L, C = wb["Ledger"], wb["Change log"]
lh = [c.value for c in L[1]]
rows = {L.cell(r, 1).value: r for r in range(2, L.max_row + 1) if L.cell(r, 1).value}
col = lambda n: lh.index(n) + 1
if "N-40" in rows:
    sys.exit("already applied; nothing written")
log = []


def note(lid, field, old, new, reason):
    log.append([D, lid, field, old, new, reason, "INT-T-012 / V5", BY])


def setf(lid, field, new, reason):
    r = rows[lid]; c = col(field); old = L.cell(r, c).value; L.cell(r, c).value = new; note(lid, field, old, new, reason)


def known(lid, text):
    r = rows[lid]; c = col("What is known now"); old = L.cell(r, c).value or ""
    L.cell(r, c).value = (old + " | " if old else "") + f"V5 {D}: " + text
    note(lid, "What is known now", "(appended)", text[:180], "design/package-v5/10 s9")


known("N-37", "third probe (run by the human, nonce-marked variants): GLM-5.3 and MiniMax-M3 are plain prefix caches (a change rewrites everything after it; the tools-first and system-first template orders differ); gpt-oss-120b is not monotonic in the edit position and unexplained; the earlier near-full cache after tools changes is consistent with earlier unrecorded runs priming those variants")
setf("N-37", "Next action", "Optional: lifetime mode after a longer gap; gpt-oss-120b's pattern is left unexplained", "design/package-v5/10 s9")
setf("N-37", "Status date", D, "INT-T-012")
known("N-31", "cache lifetime: at least about 75 minutes on gpt-oss-120b, MiniMax-M3 and GLM-5.3; the base prefix was gone after about 5 hours on gpt-oss-120b and MiniMax-M3 and still cached after about 5.5 hours on GLM-5.3 (if nobody touched it in between)")
setf("N-31", "Status date", D, "INT-T-012")
known("N-39", "the human asked for the price ranges, numbers and equations to be verified by a multi-agent literature review with scientific citations (REQ-U22-18, new row N-40)")

tmpl = rows["N-30"]; r = max(rows.values()) + 1
for cname in ("Active?", "Review flag", "Unanswered linked questions"):
    f = L.cell(tmpl, col(cname)).value
    L.cell(r, col(cname)).value = re.sub(rf"(?<![0-9$]){tmpl}(?![0-9])", str(r), f)
vals = {"LedgerID": "N-40", "Title": "Literature verification of the price ranges, the cost equations and the cache-mechanics claims of V5", "Current status": "OPEN",
        "Closes by (primary)": "LITERATURE-DIRECTED", "Cluster": "K14 Source & vendor verification", "Must precede (gate)": "G4 Before claims", "Priority": "High",
        "Owner": "Lead + agents", "Next action": "Scouts (T-20261027-118), stewards (T-119), challenger (T-120); revised 11 and new 13-literature-verification.md",
        "Kind": "Source check", "Level": "Item",
        "What is open (description)": "design/package-v5/11-price-ranges.md rests on benchmark and aggregator pages; its ranges, the cost equation of DPR-0006 and the break-even equations need primary or scientific citations (the human's instruction REQ-U22-18)",
        "What is known now": "Stored papers already print the equation structure (arXiv 2607.12161 eq. 1) and a counterfactual list-price method for open-weight models (arXiv 2608.08654); the review is in progress",
        "Closed when": "Each range and equation is cited to a primary or peer-reviewed source, or marked as an assumption or an aggregator figure; checked by stewards and a challenger",
        "Change vs baseline": "NEW (V5, 2026-10-27)", "Status date": D, "Status basis (evidence pointer)": "requirements/intake/2026-10-27-user-third-probe-done-literature-review-request.md; ops/AUTHORIZATIONS.md (AUTH-020 amendment)",
        "Needs new AUTH?": "No (AUTH-020 amendment of 2026-10-27)", "Origin (first source)": "INT-T-012 line 9", "Linked questions": ""}
for k, v in vals.items():
    L.cell(r, col(k)).value = v
note("N-40", "(new row)", "", "OPEN", "REQ-U22-18")

note_row = next(i for i in range(C.max_row, 1, -1) if str(C.cell(i, 1).value or "").startswith("Append-only"))
C.insert_rows(note_row - 1, amount=len(log))
for i, row in enumerate(log):
    for j, v in enumerate(row):
        C.cell(note_row - 1 + i, j + 1).value = v
bad = [i for i in range(2, L.max_row + 1) if L.cell(i, 1).value and any(not (isinstance(L.cell(i, c).value, str) and re.search(rf"{ref}{i}\b", L.cell(i, c).value)) for c, ref in ((4, "C"), (5, "A"), (6, "A")))]
assert not bad, f"formula check failed at rows {bad[:5]}"
wb.save(SRC)
print("change-log rows", len(log))
