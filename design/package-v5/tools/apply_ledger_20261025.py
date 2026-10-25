"""Record the human's answers of 2026-10-25 (INT-T-008) in the ledger, in place (never rebuilt): Q-U30 answered, Q-U32 carried out
(N-34 closed by evidence and moved to Archive), Q-U33 deferred by the human (N-31 next action), N-30 note. Appends Change log rows above
the append-only note. Refuses to run twice. Run from the repository root; then run design/ledger/export_ledger.py.
"""
import re, sys
from openpyxl import load_workbook

SRC = "design/ledger/PROJECT_A_OPEN_ITEMS_LEDGER.xlsx"
D, BY = "2026-10-25", "Claude"
SRCN = "INT-T-008 (requirements/intake/2026-10-25-user-answers-q-u30-expose-files.md)"
wb = load_workbook(SRC)
L, A, Q, C = wb["Ledger"], wb["Archive"], wb["Questions"], wb["Change log"]
lh = [c.value for c in L[1]]
qh = [c.value for c in Q[1]]
assert lh == [c.value for c in A[1]], "Ledger and Archive headers differ"
lrow = {L.cell(r, 1).value: r for r in range(2, L.max_row + 1) if L.cell(r, 1).value}
qrow = {Q.cell(r, 1).value: r for r in range(2, Q.max_row + 1) if Q.cell(r, 1).value}
qc = lambda name: qh.index(name) + 1
if Q.cell(qrow["Q-U30"], qc("YOUR ANSWER (type here)")).value:
    sys.exit("already applied; nothing written")
log = []


def note(lid, field, old, new, reason, src="INT-T-008 / V5"):
    log.append([D, lid, field, old, new, reason, src, BY])


# Q-U30 answered
r = qrow["Q-U30"]
for k, v in {"YOUR ANSWER (type here)": "[Answered] lines 1-2: A (total bill over total turns) as the headline, with the lead's recommendation (B, the mean of per-session rates) shown beside it",
             "Answer type": "You decided", "Answer date": D, "Review status": "REVIEWED", "Provenance (verbatim? where from)": SRCN,
             "Revised on / by": f"{D} / {BY}", "Resulting ledger changes": "see the Change log of 2026-10-25"}.items():
    Q.cell(r, qc(k)).value = v
note("Q-U30", "Answer", "", "answered (REQ-U22-08)", SRCN)

# Q-U32 carried out
r = qrow["Q-U32"]
Q.cell(r, qc("YOUR ANSWER (type here)")).value = (Q.cell(r, qc("YOUR ANSWER (type here)")).value or "") + " | Carried out 2026-10-25: both files placed in design/package-v5/inputs/expose/ and checked, ideas only (design/package-v5/09-expose-ideas-checked.md)"
Q.cell(r, qc("Review status")).value = "REVIEWED"
Q.cell(r, qc("Resulting ledger changes")).value = "see the Change log of 2026-10-25"
note("Q-U32", "Review status", "NEEDS FOLLOW-UP", "REVIEWED", "files placed and checked: design/package-v5/09")


def setf(lid, field, new, reason):
    rr = lrow[lid]; c = lh.index(field) + 1
    old = L.cell(rr, c).value; L.cell(rr, c).value = new
    note(lid, field, old, new, reason)


# N-30: primary averaging decided
rr = lrow["N-30"]; c = lh.index("What is known now") + 1
L.cell(rr, c).value = (L.cell(rr, c).value or "") + f" | V5 {D}: the headline averaging is decided: total bill over total turns, with the mean of per-session rates beside it (REQ-U22-08); DPR-0006 approval still the human's"
note("N-30", "What is known now", "(appended)", "headline averaging decided (REQ-U22-08)", "Q-U30")
setf("N-30", "Status date", D, "Q-U30")

# N-31: Q-U33 deferred by the human
setf("N-31", "Next action", "The human runs design/package-v5/tools/p0_api_probe.py with the key when able (2026-10-25: cannot now; nothing else waits on it, REQ-U22-10)", "INT-T-008 line 9")
setf("N-31", "Status date", D, "INT-T-008")

# N-34: closed by evidence, moved to Archive
setf("N-34", "Current status", "CLOSED - BY EVIDENCE", "design/package-v5/09-expose-ideas-checked.md")
setf("N-34", "Decision / result (on closure)", "The exposé's ideas (11) and the experiment brief's ideas (12) were checked as ordinary candidates against the stored evidence and V5's arithmetic: 9 kept, 8 revised, 1 parked, 5 dropped; two adopted as V5-R10 (tiered appended entries) and V5-R11 (churn, offered-set recall, oracle ablation). The brief's status notes about an earlier experiment were not used (REQ-U21-02).", "design/package-v5/09")
setf("N-34", "Status date", D, "Q-U32")
setf("N-34", "Status basis (evidence pointer)", "design/package-v5/09-expose-ideas-checked.md; INT-T-008", "Q-U32")
lid = "N-34"
lr = lrow[lid]
vals = [L.cell(lr, c).value for c in range(1, len(lh) + 1)]
vals[lh.index("Active?")] = "No"; vals[lh.index("Review flag")] = ""; vals[lh.index("Unanswered linked questions")] = 0
n = next(r for r in range(A.max_row, 1, -1) if str(A.cell(r, 1).value or "").startswith("Closed rows live here"))
A.insert_rows(n - 1, 1)
for j, v in enumerate(vals):
    A.cell(n - 1, j + 1).value = v
L.delete_rows(lr, 1)
for r in range(lr, L.max_row + 1):
    for c in range(1, len(lh) + 1):
        v = L.cell(r, c).value
        if isinstance(v, str) and v.startswith("="):
            L.cell(r, c).value = re.sub(rf"(?<![0-9$:]){r + 1}(?![0-9])", str(r), v)
note(lid, "Sheet", "Ledger", "Archive", "closed 2026-10-25", "INT-T-008 / V5")
bad = [r for r in range(2, L.max_row + 1) if L.cell(r, 1).value and any(not (isinstance(L.cell(r, c).value, str) and re.search(rf"{ref}{r}\b", L.cell(r, c).value)) for c, ref in ((4, "C"), (5, "A"), (6, "A")))]
assert not bad, f"formula re-index failed at rows {bad[:5]}"

note_row = next(r for r in range(C.max_row, 1, -1) if str(C.cell(r, 1).value or "").startswith("Append-only"))
C.insert_rows(note_row - 1, amount=len(log))
for i, row in enumerate(log):
    for j, v in enumerate(row):
        C.cell(note_row - 1 + i, j + 1).value = v
wb.save(SRC)
print(f"change-log rows {len(log)}")
