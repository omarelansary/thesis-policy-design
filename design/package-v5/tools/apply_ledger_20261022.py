"""Apply the V5 ledger changes the human accepted (Q-U31, INT-T-006 line 8) and record the answers of INT-T-006.
Edits design/ledger/PROJECT_A_OPEN_ITEMS_LEDGER.xlsx in place (never rebuilds it), appends Change log rows above the
append-only note, then the caller runs `python design/ledger/export_ledger.py`. Refuses to run twice.
Run from the repository root: python design/package-v5/tools/apply_ledger_20261022.py
Source of every change: design/package-v5/08-ledger-proposals.md (sections 1-3), updated with INT-T-006 and the P0
part-1 reading of the university API's status data.
"""
import re, sys
from openpyxl import load_workbook

SRC = "design/ledger/PROJECT_A_OPEN_ITEMS_LEDGER.xlsx"
D = "2026-10-22"
BY = "Claude"
SOURCE = "V5 (AUTH-020); accepted by the human Q-U31 (INT-T-006 line 8)"
wb = load_workbook(SRC)
L, Q, K, C = wb["Ledger"], wb["Questions"], wb["Links"], wb["Change log"]
lh = [c.value for c in L[1]]
qh = [c.value for c in Q[1]]
col = lambda name: lh.index(name) + 1
qcol = lambda name: qh.index(name) + 1
rows = {L.cell(r, 1).value: r for r in range(2, L.max_row + 1) if L.cell(r, 1).value}
qrows = {Q.cell(r, 1).value: r for r in range(2, Q.max_row + 1) if Q.cell(r, 1).value}
if "N-31" in rows or "Q-U33" in qrows:
    sys.exit("already applied; nothing written")
log = []


def note(lid, field, old, new, reason):
    log.append([D, lid, field, old, new, reason, SOURCE, BY])


def append_known(lid, text):
    r = rows[lid]; c = col("What is known now")
    old = L.cell(r, c).value or ""
    L.cell(r, c).value = (old + " | " if old else "") + f"V5 {D}: " + text
    note(lid, "What is known now", "(appended)", text[:200], "design/package-v5/08 section 1")


def set_field(lid, field, new, reason):
    r = rows[lid]; c = col(field)
    old = L.cell(r, c).value
    L.cell(r, c).value = new
    note(lid, field, old, new, reason)


V5 = "design/package-v5/08-ledger-proposals.md"
changes = [
 # id, new status or None, known text, next action, basis
 ("N-6", "PARTLY ANSWERED", "personalized tool memory exists over small tool sets (PRTA, memory-conditioned tool calling; design/package-v5/03 s5); per-user seeding of tool definitions from a large catalogue with the bill measured not found in a limited search (inputs/search-log-20261022.md s2); now HYP-0017, in the tested policy per REQ-J22-01", "05 P1 (needs AUTH)", "design/package-v5/03 s5; search log s2; INT-T-006 line 23"),
 ("N-7", "PARTLY ANSWERED", "resolved in design: per-user seeds placed after a shared, task-independent prefix (04 V5-R06); provider sharing scope still unstated; prefix caching on the university API unknown (05 P0 part 2)", "05 P0 part 2", "design/package-v5/04 V5-R06"),
 ("AS-07", None, "Executor README read (UsefulSoftwareCo/executor, MIT, 4,235 stars, 2026-10-22); its MCP surface is not in the README", "Read Executor's MCP docs (05 P0); Q-J33 reworded", "L/v5-executor-readme-20261022.md:5-7, :123"),
 ("C-D08/EXEC", None, "identified as above; required baseline (REQ-O02-06); Q-J33 reworded, assumption 'as shipped'", "Q-J33", "design/package-v5/05 s4"),
 ("AS-09", None, "vendors ship partial-information tiers: an OpenAI deferred function shows name and description, schema deferred; a namespace shows only its own description", "Q-J28 (source the supervisor meant)", "L/v5-openai-tool-search-20261022.md:22"),
 ("AS-13", "PARTLY ANSWERED", "the tier form exists in vendor products; its accuracy effect is not reported; REQ-J22-02: count reduction was an example, the goal is the bill", "Graded-exposure arm only if the loop needs it", "L/v5-openai-tool-search-20261022.md:22; INT-T-006 line 25"),
 ("AS-15", None, "desk arithmetic done (D-02 part 2): consistent within one conversation against re-sending; grows with every request against not adding; his words INT-T-002 L209, L211, L219, L221 (design/package-v5/02 (b))", "Billing reading 05 P0 part 2", "design/package-v5/calc/memory_seed.stdout.txt; design/package-v5/02"),
 ("C-S04", None, "TTL and gap cost now in D-01 and D-02 revision 3", "Our gap distribution (05 P1)", "design/package-v5/calc/"),
 ("C-S02/INVERSION", None, "D-01 locates the regime: break-even about 10-230 tools, by searches per session, definition length and whether the small core caches", "05 P5", "design/package-v5/calc/regime_bill.stdout.txt"),
 ("C-S03/DEFUSAGE", None, "OpenAI deferred functions still send name and description; on the university API vendor deferral is not available (open-weight models, 05 P0 part 1), so a client search applies", "05 P0 part 2, P5", "L/v5-openai-tool-search-20261022.md:22; L/v5-scads-llm-state-20261022.json"),
 ("C-S11/TRIG", "DEFERRED", "dropped from V5-R's default: no source supports a system trigger (04 V5-R07)", None, "design/package-v5/04 V5-R07"),
 ("C-S11/CASCADE", "PARTLY ANSWERED", "hierarchical retrieval exists (AnyTool); Copilot ships two levels and states a cache miss per group expansion (EXT-0228); no multi-stage run-time policy measured on the bill found in a limited search", None, "L/v5-anytool-2402.04253.txt:16-37; EXT-0228"),
 ("C-E02/E31", "PARTLY ANSWERED", "ScaleMCP, MemTool, Manus, langgraph-bigtool stored and checked by stewards (FND-0040..0042); Spring AI and Toolshed remain unread", "Spring AI (origin r4-spring-tool-search-blog), Toolshed", "FND-0040, FND-0041, FND-0042"),
 ("C-E02/E25", None, "one cache-bust issue stored with community data; Claude Code's tool-search docs stored", None, "L/v5-ccissue-92033-cachebust-20261022.txt:180-192; L/v5-claudecode-mcp-20261022.md:1441"),
 ("C-D08/ARMS-CACHE", "PARTLY ANSWERED", "arms fixed for the loop: DPR-0008 against V5-R on returning sessions (05 P1), all-offered against subset-offered (05 P2)", "05 P1, P2", "design/package-v5/05 s2"),
 ("C-S08/UNINSTR", None, "now HYP-0019 with a test design", "05 P3", "design/package-v5/05 P3"),
 ("C-X07", None, "split into P0 (API fields; part 1 done from the status page) and P5 (regime reading)", "05 P0 part 2, P5", "design/package-v5/05 s2"),
]
for lid, status, known, nxt, basis in changes:
    if lid not in rows:
        sys.exit(f"ledger row {lid} not found; nothing written")
for lid, status, known, nxt, basis in changes:
    if status:
        set_field(lid, "Current status", status, f"{V5} section 1; basis {basis}")
    append_known(lid, known)
    if nxt:
        set_field(lid, "Next action", nxt, f"{V5} section 1")
    if lid == "C-S11/TRIG":
        set_field(lid, "Reopen trigger", "05 P3 shows that models rarely search when the needed tool is deferred", f"{V5} section 1")
    set_field(lid, "Status date", D, "V5")
    set_field(lid, "Status basis (evidence pointer)", basis, "V5")

# New ledger rows, formulas copied from an existing row with the row number adjusted
tmpl = rows["N-30"]
last = max(rows.values())
new = [
 ("N-31", "University API: cache fields and which tool mechanisms pass through (HYP-0020)", "PARTLY ANSWERED", "RECORDS-AT-HAND", "K9 Cache & billing", "Scientific unknown", "Whether the university API reports cache writes and reads and honours allowed_tools", "P0 part 1 (2026-10-22): open-weight chat models with tool support, embedding and rerank models, endpoint /v1/ (L/v5-scads-llm-state-20261022.json; L/v5-tud-selfservice-scads-llm-20261022.txt:46); no Anthropic or OpenAI models, so the client-side route applies", "One keyed request pair per model read", "05 P0 part 2 (Q-U33)", "Q-U33", "Yes (Q-U33)", "design/package-v5/05 P0"),
 ("N-32", "Offering fewer tools at a fixed cached prefix: accuracy effect (HYP-0018)", "OPEN", "TEST-PILOT", "K3 Policy form", "Test", "Accuracy when a subset is offered while all definitions stay cached", "Vendor offering mechanisms are not available on the university API; whether it honours allowed_tools is unknown (N-31)", "05 P2 result", "05 P2 after N-31", "", "Pilot: yes", "design/package-v5/05 P2"),
 ("N-33", "Regime switch by D-01: does the arithmetic hold on our model and catalogue?", "OPEN", "TEST-PILOT", "K9 Cache & billing", "Test", "Where the break-even catalogue size lies on the route used", "D-01 revision 2 gives about 10-230 tools under stated assumptions", "05 P5 reading", "05 P5", "", "Pilot: yes", "design/package-v5/calc/regime_bill.stdout.txt"),
 ("N-34", "Exposé ideas to be tested against V5's result (Q-U32)", "OPEN", "USER-DECISION", "K6 Deliverable & authorization", "Source check", "Check the June exposé's ideas only, as ordinary candidates (INT-T-006 line 10)", "The exposé and the experiment brief named by the human are outside the workspace; the guard blocks reading them; copies inside the workspace are needed", "Each idea kept, revised or dropped with a reason", "The human copies the exposé into the workspace", "Q-U32", "No (AUTH-020 limits)", "INT-T-006 lines 10-20"),
 ("N-35", "The lead's git push is refused by the auto-mode classifier even with allow rules", "OPEN", "USER-DECISION", "K16 Limits & provenance", "Safeguard limit", "Pushes of v5-audit are done by the human", "Observed in WD-042 preflight", "Accepted as a limitation or a settings change by the human", "Human decides", "", "No", "ops/DESIGN-DECISIONS.md WD-042"),
 ("N-36", "V5 rows are model work checked by model work; no human spot-check yet", "OPEN", "USER-DECISION", "K16 Limits & provenance", "Limitation", "A human check of a sample of design/package-v5/03 rows", "Stewards, challengers and a rebuttal instance checked rows; no human check", "Human spot-check done or limitation accepted", "Human decides", "", "No", "design/package-v5/00-README.md"),
]
fixed = {"Kind", "Level", "Change vs baseline", "Priority", "Owner"}
for i, (lid, title, status, closes, cluster, kind, open_, known, closed_when, nxt, links, auth, basis) in enumerate(new):
    r = last + 1 + i
    for cname in ("Active?", "Review flag", "Unanswered linked questions"):
        f = L.cell(tmpl, col(cname)).value
        L.cell(r, col(cname)).value = re.sub(rf"(?<![0-9$]){tmpl}(?![0-9])", str(r), f)
    vals = {"LedgerID": lid, "Title": title, "Current status": status, "Closes by (primary)": closes, "Cluster": cluster,
            "Must precede (gate)": "G2 Before the affected pilot/arm", "Priority": "Normal", "Owner": "Lead; you decide",
            "Next action": nxt, "Kind": kind, "Level": "Item", "What is open (description)": open_, "What is known now": known,
            "Closed when": closed_when, "Linked questions": links, "Change vs baseline": "NEW (V5, 2026-10-22)",
            "Status date": D, "Status basis (evidence pointer)": basis, "Needs new AUTH?": auth,
            "Origin (first source)": "design/package-v5/08-ledger-proposals.md section 2"}
    for k, v in vals.items():
        L.cell(r, col(k)).value = v
    note(lid, "(new row)", "", status, f"{V5} section 2")

# Questions: new rows with the answers of INT-T-006
qtmpl = qrows["Q-J30"]; qlast = max(qrows.values())
prov = "INT-T-006 (requirements/intake/2026-10-22-user-answers-after-v5.md)"
qs = [
 ("Q-J31", "Joaa", "Is the staged form satisfied by three tool-deciding stages (deployment configuration, memory-informed offering, discovery with a ranker)?", "Yes | No | Other", "Policy form of V5-R", "design/package-v5/05 s4", "[Answered by the student on Joaa's behalf] line 23: a starting point; the staged form was an example and may change; memory must not stay off, test it properly in the right place", "You decided", D, "REVIEWED"),
 ("Q-J32", "Joaa", "Does 'reduce the count of tools' mean tools offered or definitions billed?", "Offered | Billed | Other", "Objective", "design/package-v5/05 s4", "[Answered by the student on Joaa's behalf] line 25: the count was an example of what could lower the bill; focus on the goal", "You decided", D, "REVIEWED"),
 ("Q-J33", "Joaa", "Executor baseline: as its authors ship it, or a particular setup he has in mind?", "As shipped | A setup: ... | No preference", "Defines the Executor arm", "design/package-v5/05 s4 (reworded)", "", "", "", ""),
 ("Q-U28", "You", "Authorise P0 (university API reading, Executor docs)?", "Yes | No", "First loop item", "design/package-v5/05 s5", "[Answered] line 2: fetch the status page (done); no keyed request authorised by these words", "You decided", D, "NEEDS FOLLOW-UP"),
 ("Q-U29", "You", "Take DPR-0007 as the working draft?", "Yes | Keep V4 | Wait", "Which draft goes to the supervisor", "design/package-v5/05 s5", "[Answered] line 4: yes, working draft", "You decided", D, "REVIEWED"),
 ("Q-U30", "You", "Which average of the bill is primary: total over turns, or mean of per-session rates?", "Total over turns | Mean of session rates | Both", "Reading of every result", "design/package-v5/05 s5 (explained with an example)", "", "", "", ""),
 ("Q-U31", "You", "Apply the ledger changes proposed in design/package-v5/08?", "Yes | No | Some: ...", "Ledger state", "design/package-v5/05 s5", "[Answered] line 8: accepted (applied 2026-10-22)", "You decided", D, "REVIEWED"),
 ("Q-U32", "You", "Supply the June exposé so its ideas can be tested against V5?", "Yes | Later | No", "Exposé ideas (N-34)", "design/package-v5/05 s5", "[Answered] line 10: check its ideas only, as ordinary candidates; files are outside the workspace", "You decided", D, "NEEDS FOLLOW-UP"),
 ("Q-U33", "You", "Authorise P0 part 2: a few requests with your key to read usage records and test allowed_tools?", "Yes | No", "Cache terms on the university API (N-31)", "design/package-v5/05 s5", "", "", "", ""),
]
for i, (qid, who, text, fmt, why, aka, ans, atype, adate, rstat) in enumerate(qs):
    r = qlast + 1 + i
    f = Q.cell(qtmpl, qcol("Ledger items open because of it (formula)")).value
    Q.cell(r, qcol("Ledger items open because of it (formula)")).value = re.sub(rf"(?<![0-9$]){qtmpl}(?![0-9])", str(r), f)
    for k, v in {"QID": qid, "Asked of": who, "Question (neutral wording; one issue each)": text, "Answer format / options": fmt,
                 "Why it matters": why, "Also known as": aka, "Ask at step": "S05", "YOUR ANSWER (type here)": ans or None,
                 "Answer type": atype or None, "Answer date": adate or None, "Review status": rstat or None,
                 "Provenance (verbatim? where from)": prov if ans else None, "Revised on / by": f"{D} / {BY}" if ans else None,
                 "Resulting ledger changes": "see the Change log of 2026-10-22" if ans else None}.items():
        Q.cell(r, qcol(k)).value = v
    note(qid, "(new question)", "", "answered" if ans else "open", "design/package-v5/05 sections 4-5; INT-T-006")

# Links
klast = max(r for r in range(2, K.max_row + 1) if K.cell(r, 1).value)
ktmpl = klast
links = [("Q-U33", "N-31", "CLOSES"), ("Q-U32", "N-34", "CLOSES"), ("Q-J33", "C-D08/EXEC", "NARROWS"), ("Q-U30", "N-30", "NARROWS"),
         ("Q-J31", "N-6", "INFORMS"), ("Q-J32", "C-D02", "NARROWS"), ("Q-U29", "N-16", "INFORMS")]
for i, (qid, lid, eff) in enumerate(links):
    r = klast + 1 + i
    K.cell(r, 1).value, K.cell(r, 2).value, K.cell(r, 3).value = qid, lid, eff
    for c in (4, 5):
        K.cell(r, c).value = re.sub(rf"(?<![0-9$]){ktmpl}(?![0-9])", str(r), K.cell(ktmpl, c).value)
    note(lid, "Link", "", f"{qid} {eff}", "design/package-v5/05")

# Change log: insert above the blank row and the append-only note
note_row = next(r for r in range(C.max_row, 1, -1) if str(C.cell(r, 1).value or "").startswith("Append-only"))
ins = note_row - 1
C.insert_rows(ins, amount=len(log))
for i, row in enumerate(log):
    for j, v in enumerate(row):
        C.cell(ins + i, j + 1).value = v
wb.save(SRC)
print(f"ledger rows changed {len(changes)}, new rows {len(new)}, questions {len(qs)}, links {len(links)}, change-log rows {len(log)}")
