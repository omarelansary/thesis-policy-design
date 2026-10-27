"""Record the second API probe (INT-T-010, REQ-U22-12, REQ-U22-13) in the ledger, in place (never rebuilt): N-37 and N-31 updated, C-S08
note, new rows N-38 (which models) and N-39 (named price table). Appends Change log rows above the append-only note. Refuses to run twice.
Run from the repository root, then run design/ledger/export_ledger.py.
"""
import re, sys
from openpyxl import load_workbook

SRC = "design/ledger/PROJECT_A_OPEN_ITEMS_LEDGER.xlsx"
D, BY = "2026-10-27", "Claude"
SRCN = "INT-T-010 (requirements/intake/2026-10-27-user-report-second-probe-model-preference.md); design/package-v5/10-p0-university-api-reading.md s7"
wb = load_workbook(SRC)
L, C = wb["Ledger"], wb["Change log"]
lh = [c.value for c in L[1]]
rows = {L.cell(r, 1).value: r for r in range(2, L.max_row + 1) if L.cell(r, 1).value}
if "N-38" in rows:
    sys.exit("already applied; nothing written")
col = lambda n: lh.index(n) + 1
log = []


def note(lid, field, old, new, reason):
    log.append([D, lid, field, old, new, reason, "INT-T-010 / V5", BY])


def setf(lid, field, new, reason):
    r = rows[lid]; c = col(field); old = L.cell(r, c).value; L.cell(r, c).value = new; note(lid, field, old, new, reason)


def known(lid, text):
    r = rows[lid]; c = col("What is known now"); old = L.cell(r, c).value or ""
    L.cell(r, c).value = (old + " | " if old else "") + f"V5 {D}: " + text
    note(lid, "What is known now", "(appended)", text[:180], "design/package-v5/10 s7")


known("N-37", "second probe (3 models): tool calling works (auto, required, named; 9 of 9); an absent tool: gpt-oss and MiniMax-M3 declined, GLM-5.3 called a wrong tool (tool_09); a changed tools array rewrote 65% (added tool) to 99% (first tool edited) of the prompt on GLM-5.3 (supports V5-R02) but left gpt-oss and MiniMax-M3 almost fully cached, unexplained (an earlier unrecorded run, or a non-prefix cache)")
setf("N-37", "Current status", "PARTLY ANSWERED", "design/package-v5/10 s7")
setf("N-37", "Next action", "Third probe: the human runs design/package-v5/tools/p0_api_probe_v3.py (nonce-marked variants; lifetime mode later; GLM-5.3 with --pause 35)", "design/package-v5/10 s9")
setf("N-37", "Status date", D, "INT-T-010")
known("N-31", "cache lived at least about 75 minutes on gpt-oss, MiniMax-M3 and GLM-5.3 (first request of the second probe was already cached from the first probe, if no run intervened); MiniMax-M3 re-creates a 128-token block on every request; no prices")
setf("N-31", "Status date", D, "INT-T-010")
known("C-S08", "one absent-tool request per model through the university API: 2 of 3 models declined, GLM-5.3 made a wrong-tool call (single runs; design/package-v5/10 s7)")
setf("C-S08", "Status date", D, "INT-T-010")

tmpl = rows["N-30"]; last = max(rows.values())
new = [
 ("N-38", "Which models the proof-of-concept loop uses (REQ-U21-08); the human leans away from gpt-oss-120b", "OPEN", "USER-DECISION", "K5 Evaluation contract", "Decision",
  "Choose the models for the loop and the end-to-end study from those the university API serves with tool support",
  "Seven open-weight chat models report tool support (status data 2026-10-22); three were probed twice: all call tools correctly in the simple cases, caching and rate limits differ (GLM-5.3: 2 requests, clean prefix cache). The human prefers models other than gpt-oss-120b and calls this a personal preference (REQ-U22-13). The data shows no model to exclude; the end-to-end study needs at least two families to read transfer (AS-12)",
  "A decision recorded in the words of the human, with the reason", "The human chooses", "design/package-v5/10 s8; REQ-U22-13"),
 ("N-39", "A named price table for the bill on the university route (the API is free and reports tokens only)", "OPEN", "USER-DECISION", "K2 Objective & accounting", "Parameter",
  "DPR-0006 needs one price per component and per model; the university API sends no bill, so a table must be named (for example a hosted list price for a comparable model) or a symbolic price ratio reported",
  "The API reports cached, created and reasoning tokens but no prices (probes 2026-10-27). Whether cache writes carry a premium on this service is not stated; MiniMax-M3 writes a 128-token block on every request",
  "The human names the table or chooses to report token classes and price ratios only", "The human decides", "DPR-0006 s1; design/package-v5/10 s2"),
]
for i, (lid, title, status, closes, cluster, kind, open_, known_, closed_when, nxt, basis) in enumerate(new):
    r = last + 1 + i
    for cname in ("Active?", "Review flag", "Unanswered linked questions"):
        f = L.cell(tmpl, col(cname)).value
        L.cell(r, col(cname)).value = re.sub(rf"(?<![0-9$]){tmpl}(?![0-9])", str(r), f)
    vals = {"LedgerID": lid, "Title": title, "Current status": status, "Closes by (primary)": closes, "Cluster": cluster,
            "Must precede (gate)": "G1 Before any implementation", "Priority": "High", "Owner": "You decide; lead prepares", "Next action": nxt,
            "Kind": kind, "Level": "Item", "What is open (description)": open_, "What is known now": known_, "Closed when": closed_when,
            "Change vs baseline": "NEW (V5, 2026-10-27)", "Status date": D, "Status basis (evidence pointer)": basis,
            "Needs new AUTH?": "No", "Origin (first source)": "design/package-v5/10", "Linked questions": ""}
    for k, v in vals.items():
        L.cell(r, col(k)).value = v
    note(lid, "(new row)", "", status, "design/package-v5/10")

note_row = next(i for i in range(C.max_row, 1, -1) if str(C.cell(i, 1).value or "").startswith("Append-only"))
C.insert_rows(note_row - 1, amount=len(log))
for i, row in enumerate(log):
    for j, v in enumerate(row):
        C.cell(note_row - 1 + i, j + 1).value = v
bad = [i for i in range(2, L.max_row + 1) if L.cell(i, 1).value and any(not (isinstance(L.cell(i, c).value, str) and re.search(rf"{ref}{i}\b", L.cell(i, c).value)) for c, ref in ((4, "C"), (5, "A"), (6, "A")))]
assert not bad, f"formula check failed at rows {bad[:5]}"
wb.save(SRC)
print("change-log rows", len(log))
